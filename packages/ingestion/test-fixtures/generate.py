"""SYNTHETIC only. /usr/bin/python3 -B generate.py. Requires openpyxl 3.1.5.
ZIP timestamps normalized for byte-reproducible artifacts; no customer data.
"""
from pathlib import Path
from io import BytesIO
from datetime import datetime
import hashlib, json, zipfile, re
import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.worksheet.filters import FilterColumn, Filters, CustomFilters, CustomFilter
from openpyxl.workbook.defined_name import DefinedName
assert openpyxl.__version__ == '3.1.5'
root = Path(__file__).parent
manifest = {}
def save(name, wb):
    wb.properties.created = wb.properties.modified = datetime(2020, 1, 1)
    b = BytesIO(); wb.save(b)
    with zipfile.ZipFile(b) as src:
        parts = {n: src.read(n) for n in src.namelist()}
        parts['docProps/core.xml'] = re.sub(rb'(<dcterms:modified[^>]*>)[^<]+', rb'\g<1>2020-01-01T00:00:00Z', parts['docProps/core.xml'])
        write(name, parts)
def write(name, parts):
    dest = root / (name + '.xlsx')
    with zipfile.ZipFile(dest, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, data in sorted(parts.items()):
            info = zipfile.ZipInfo(n, (2020,1,1,0,0,0)); info.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(info, data)
    manifest[name] = hashlib.sha256(dest.read_bytes()).hexdigest()
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Mensajes'
ws.append(['message', 'amount']); ws.append(['SYNTHETIC: café 😀\nsegunda línea', 12.34])
save('openpyxl-basic', wb)
ws['A3'] = CellRichText('inicio ', TextBlock(InlineFont(b=True), '尾😀'), ' fin')
ws['B3'] = '=SUM(B2,1)'
ws2 = wb.create_sheet("Ventas 'México'"); ws2.append(['texto', 'número']); ws2.append(['SYNTHETIC segunda hoja', -99]); ws2.print_area = 'A1:B2'
save('openpyxl-multiple-rich-formula', wb)
ws.auto_filter.ref = 'A1:B3'
ws.auto_filter.filterColumn.append(FilterColumn(colId=0, filters=Filters(filter=['SYNTHETIC: café 😀\nsegunda línea'])))
ws.auto_filter.add_sort_condition('B2:B3')
ws.print_area = 'A1:B3'; ws.print_title_rows = '1:1'; ws.print_title_cols = 'A:A'
ws.freeze_panes = 'A2'; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.orientation = 'landscape'; ws.oddHeader.center.text = 'SYNTHETIC &P'
wb.security.lockStructure = True; wb.security.set_workbook_password('synthetic')
wb.defined_names.add(DefinedName('SyntheticRange', attr_text="'Mensajes'!$B$2"))
save('openpyxl-metadata', wb)
ws.auto_filter.filterColumn = [FilterColumn(colId=1, customFilters=CustomFilters(customFilter=[CustomFilter(operator='greaterThan', val='0')]))]
save('openpyxl-custom-filter', wb)
ws.row_dimensions[2].hidden = True; ws.column_dimensions['B'].hidden = True
wb.defined_names.add(DefinedName('SyntheticFormula', attr_text='SUM(Mensajes!$B$2,1)'))
save('openpyxl-hidden-and-named-formula', wb)
with zipfile.ZipFile(root/'openpyxl-metadata.xlsx') as z: base = {n:z.read(n) for n in z.namelist()}
def bad(name, part, old, new):
    parts = base.copy(); text = parts[part].decode(); assert old in text, (name, old)
    parts[part] = text.replace(old, new, 1).encode(); write(name, parts)
bad('bad-protection-child','xl/workbook.xml','<workbookProtection ', '<workbookProtection><sheets/></workbookProtection><workbookProtection ')
bad('bad-name-child','xl/workbook.xml','</definedName>', '<row r="9"/></definedName>')
bad('bad-name-scope','xl/workbook.xml','localSheetId="0"','localSheetId="99"')
bad('bad-name-external','xl/workbook.xml',"'Mensajes'!$B$2", "'[external.xlsx]Mensajes'!$B$2")
bad('bad-filter-child','xl/worksheets/sheet1.xml','<filters>', '<filters><row r="9"/>')
bad('bad-filter-column','xl/worksheets/sheet1.xml','colId="0"','colId="99"')
bad('bad-filter-range','xl/worksheets/sheet1.xml','<autoFilter ref="A1:B3"','<autoFilter ref="B3:A1"')
bad('bad-protection-boolean','xl/workbook.xml','lockStructure="1"','lockStructure="perhaps"')
bad('bad-header-cell','xl/worksheets/sheet1.xml','</oddHeader>', '<c r="A9"><v>99</v></c></oddHeader>')
bad('bad-duplicate-names','xl/workbook.xml','</definedNames>', '<definedName name="SyntheticRange">Mensajes!A1</definedName></definedNames>')
bad('bad-duplicate-protection','xl/workbook.xml','</sheets>', '</sheets><workbookProtection/>')
bad('bad-name-attribute','xl/workbook.xml','name="SyntheticRange"', 'name="SyntheticRange" unknown="1"')
bad('bad-name-macro','xl/workbook.xml','name="SyntheticRange"', 'name="SyntheticRange" xlm="1"')
bad('bad-print-area','xl/workbook.xml', "'Mensajes'!$A$1:$B$3", "'Mensajes'!$B$3:$A$1")
bad('bad-print-sheet','xl/workbook.xml', "'Mensajes'!$A$1:$B$3", "'Missing'!$A$1:$B$3")
bad('bad-sort-range','xl/worksheets/sheet1.xml','<sortCondition descending="0" ref="B2:B3"', '<sortCondition descending="0" ref="C2:C3"')
bad('bad-filter-choice','xl/worksheets/sheet1.xml','</filters>', '</filters><top10 val="10"/>')
bad('bad-filter-duplicate','xl/worksheets/sheet1.xml','</filterColumn>', '</filterColumn><filterColumn colId="0"/>')
bad('bad-protection-hash-group','xl/workbook.xml','lockStructure="1"', 'lockStructure="1" workbookHashValue="YQ=="')
(root/'manifest.json').write_text(json.dumps({'label':'SYNTHETIC','writer':'openpyxl 3.1.5','generator':'generate.py','files_sha256':manifest},indent=2)+'\n')
