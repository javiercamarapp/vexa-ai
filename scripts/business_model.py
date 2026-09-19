"""Reproducible scenario model, NOT a prediction or observed business valuation."""
import csv,json
from decimal import Decimal as D
from pathlib import Path
ROOT=Path(__file__).resolve().parents[1];OUT=ROOT/'negocio/05-Precios-y-Finanzas'
def dec(v):return v if isinstance(v,D) else D(str(v))
def load():return json.loads((OUT/'supuestos.json').read_text(),parse_float=D,parse_int=D)
def sized(config,s):
 core=dec(config['core_firms']);price=dec(s['monthly_price']);eligible=core
 stages=[]
 for k in ['physical_category_fit','volume_and_need_fit','data_stack_fit','governance_fit']:
  f=dec(s[k])
  if not 0<=f<=1:raise ValueError('eligibility outside [0,1]')
  eligible*=f;stages.append({'filter':k,'conditional_assumption':f,'remaining_expected_accounts':eligible})
 return {'core_firms_observed_2022':core,'annual_price_assumed':price*12,'tam_core_price_scenario':core*price*12,'sam_expected_accounts_scenario':eligible,'sam_arr_scenario':eligible*price*12,'stages':stages}
def monthly_cost(config,price):
 c=config['costs'];n=dec(c['conversations_per_active_month']);it=dec(c['input_tokens_per_conversation']);ot=dec(c['output_tokens_per_conversation'])
 extraction=n*(it*dec(c['extraction_input_usd_per_million'])+ot*dec(c['extraction_output_usd_per_million']))/D(1000000)
 escalation=n*dec(c['escalation_fraction'])*(it*dec(c['escalation_input_usd_per_million'])+ot*dec(c['escalation_output_usd_per_million']))/D(1000000)
 ai=(extraction+escalation)*dec(c['retry_cost_multiplier'])
 support=dec(c['variable_support_hours_per_active_month'])*dec(c['support_hourly_loaded_usd'])
 cost=ai+support+dec(c['embeddings_briefs_per_active_month'])+dec(c['allocated_cloud_per_active_month'])+dec(price)*dec(c['payment_rate'])
 return {'extraction':extraction,'escalation':escalation,'ai_with_retry_allowance':ai,'support':support,'recurring_cogs_per_active_month':cost,'gross_margin':(dec(price)-cost)/dec(price)}
def forecast(config,s):
 size=sized(config,s);c=config['costs'];f=config['forecast'];price=dec(s['monthly_price']);cost=monthly_cost(config,price)['recurring_cogs_per_active_month']
 active=D(0);backlog=D(0);prospected=D(0);cash=dec(f['initial_cash']);wins=[];result=[]
 for m in range(1,int(f['months'])+1):
  y=(m-1)//12+1
  outbound=D(0) if m<int(f['first_outreach_month']) else min(dec(s[f'outreach_year{y}_monthly']),max(D(0),dec(config['core_firms'])-prospected))
  prospected+=outbound;meetings=outbound*dec(s['meeting_rate']);pilots=meetings*dec(s['pilot_rate']);won=pilots*dec(s['paid_rate']);wins.append(won)
  lag=int(s['sales_lag_months']);matured=wins[m-lag-1] if m>lag else D(0)
  lost=active*dec(s['monthly_logo_churn']);surviving=active-lost
  available=backlog+matured
  new=min(available,dec(s[f'onboarding_cap_year{y}']),max(D(0),size['sam_expected_accounts_scenario']-surviving))
  backlog=available-new;active=surviving+new
  billed=surviving+new/2 # new logos begin at mid-month; churn exits at month start
  recurring=billed*price;setup=new*dec(c['setup_fee_usd']);cogs=billed*cost
  onboarding=new*dec(c['onboarding_hours_per_new_customer'])*dec(c['onboarding_hourly_usd'])
  failed_pilots=pilots*(1-dec(s['paid_rate']))*dec(c['unconverted_pilot_cost_usd'])
  fixed=dec(s[f'fixed_opex_year{y}_monthly']);marketing=dec(c['extra_fixed_marketing_monthly'])
  net=recurring+setup-cogs-onboarding-failed_pilots-fixed-marketing;cash+=net
  result.append({'month':m,'year':y,'new_accounts_targeted':outbound,'cumulative_accounts_targeted':prospected,'meetings_expected':meetings,'pilots_expected':pilots,'wins_before_lag':won,'matured_wins':matured,'backlog_eom':backlog,'lost_logos':lost,'new_logos':new,'active_logos_eom':active,'recurring_revenue':recurring,'setup_revenue':setup,'recurring_cogs':cogs,'onboarding_cost':onboarding,'failed_pilot_cost':failed_pilots,'fixed_opex':fixed,'extra_marketing':marketing,'net_cash_flow_before_financing':net,'cash_before_financing':cash,'arr_exit':active*price*12})
 return result

def serial(x):
 if isinstance(x,D):return float(x)
 raise TypeError(type(x).__name__)
def run():
 config=load();rowsall=[];summary={}
 for name,s in config['scenarios'].items():
  rows=forecast(config,s);size=sized(config,s);unit=monthly_cost(config,s['monthly_price']);annual=[]
  for y in range(1,4):
   yr=[r for r in rows if r['year']==y];new=sum(r['new_logos'] for r in yr)
   sm=sum(r['fixed_opex']*dec(config['forecast']['sales_marketing_share_fixed_opex'])+r['extra_marketing']+r['failed_pilot_cost'] for r in yr)
   cac=sm/new if new else None;grossprofit=dec(s['monthly_price'])-unit['recurring_cogs_per_active_month']
   annual.append({'year':y,'new_expected_logos':new,'end_expected_logos':yr[-1]['active_logos_eom'],'revenue_recurring':sum(r['recurring_revenue'] for r in yr),'arr_exit':yr[-1]['arr_exit'],'net_cash_flow':sum(r['net_cash_flow_before_financing'] for r in yr),'cash_before_financing_end':yr[-1]['cash_before_financing'],'fully_allocated_cac_scenario':cac,'cac_payback_months_scenario':cac/grossprofit if cac is not None and grossprofit>0 else None})
  funding={str(n):max(D(0),-min(r['cash_before_financing'] for r in rows[:n])) for n in [12,18,24,36]}
  summary[name]={'size':size,'unit_cost':unit,'annual':annual,'cash_needed_no_buffer':funding,'cash_18_months_plus_6_month_fixed_buffer':funding['18']+6*(dec(s['fixed_opex_year2_monthly'])+dec(config['costs']['extra_fixed_marketing_monthly'])),'first_positive_month_before_financing':next((r['month'] for r in rows if r['net_cash_flow_before_financing']>=0),None),'caution':'All SOM/price/eligibility/cost/conversion values are scenarios, not measured results; positive cash flow can be temporary.'}
  rowsall.extend({'scenario':name,**r} for r in rows)
 (OUT/'resultados-modelo.json').write_text(json.dumps(summary,default=serial,indent=2,ensure_ascii=False)+'\n')
 with (OUT/'forecast-36-meses.csv').open('w') as f:
  w=csv.DictWriter(f,fieldnames=list(rowsall[0]));w.writeheader()
  for row in rowsall:w.writerow({k:format(v,'.4f') if isinstance(v,D) else v for k,v in row.items()})
 import openpyxl
 from openpyxl.styles import Font,PatternFill,Alignment
 wb=openpyxl.Workbook();ws=wb.active;ws.title='LEEME'
 for row in [['VEXA — escenarios, NO resultados'],['Fecha','2026-09-19'],['Fuente modelo','scripts/business_model.py; supuestos.json'],['IMPORTANTE','Valores calculados; cambiar JSON y regenerar. No simulador Excel con inputs enlazados.'],['Censo','2022, NAICS 2017; USD10M <= receipts < USD100M. No contar NAICS como empresas únicas.'],['Unidades','USD; clientes fraccionarios son expectativas matemáticas, no clientes reales.'],['Limitaciones','Sin impuestos, DSO ni financiación; supuestos de conversión/precio/churn no validados.']]:ws.append(row)
 ws=wb.create_sheet('Forecast36')
 ws.append(list(rowsall[0]))
 for row in rowsall:ws.append([float(v) if isinstance(v,D) else v for v in row.values()])
 ws=wb.create_sheet('Resumen')
 ws.append(['escenario','TAM núcleo precio supuesto','SAM empresas esperado','SAM ARR supuesto','Año','ARR salida','Ingresos año','Clientes equivalentes al cierre','Caja año sin financiamiento','CAC escenario','Payback meses'])
 for name,d in summary.items():
  for a in d['annual']:ws.append([name,float(d['size']['tam_core_price_scenario']),float(d['size']['sam_expected_accounts_scenario']),float(d['size']['sam_arr_scenario']),a['year'],float(a['arr_exit']),float(a['revenue_recurring']),float(a['end_expected_logos']),float(a['net_cash_flow']),float(a['fully_allocated_cac_scenario']) if a['fully_allocated_cac_scenario'] else None,float(a['cac_payback_months_scenario']) if a['cac_payback_months_scenario'] else None])
 ws=wb.create_sheet('CensusFilas')
 with (OUT/'census-filas.csv').open() as f:
  for row in csv.reader(f):ws.append(row)
 ws=wb.create_sheet('SupuestosJSON')
 for line in (OUT/'supuestos.json').read_text().splitlines():ws.append([line])
 for ws in wb:
  ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='222222')
  for col in ws.columns:
   ws.column_dimensions[col[0].column_letter].width=min(65,max(18,max(len(str(c.value or '')) for c in col[:25])+2))
 wb.save(OUT/'VEXA-MERCADO-Y-FINANZAS.xlsx')
 import copy
 sensitivity=[]
 for n in [10000,50000,150000,500000]:
  scenario=copy.deepcopy(config);scenario['costs']['conversations_per_active_month']=n
  u=monthly_cost(scenario,1499)
  sensitivity.append({'conversations_per_month':n,'MRR_assumed':1499,'COGS_scenario':round(float(u['recurring_cogs_per_active_month']),2),'gross_margin_pct':round(float(u['gross_margin'])*100,2)})
 with (OUT/'sensibilidad-volumen.csv').open('w') as f:
  writer=csv.DictWriter(f,fieldnames=list(sensitivity[0]));writer.writeheader();writer.writerows(sensitivity)
 print(json.dumps(summary,default=serial,indent=2))
if __name__=='__main__':run()
