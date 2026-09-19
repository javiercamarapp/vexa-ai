import importlib.util,json,sys,unittest
from decimal import Decimal as D
from pathlib import Path
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'scripts'))
import business_model as m
class BusinessModelTests(unittest.TestCase):
 def setUp(self):self.c=m.load();self.s=self.c['scenarios']['base']
 def test_official_core_has_all_nine_bands_and_2975_firms(self):
  data=json.loads((ROOT/'negocio/05-Precios-y-Finanzas/universo-census.json').read_text())
  core=next(x for x in data['sectors'] if x['naics']=='454110')
  self.assertEqual(core['published_firm_sum'],2975);self.assertTrue(core['complete_target_bands']);self.assertEqual(core['present_bands'],9)
  self.assertEqual(core['rows'],list(range(18333,18342)))
 def test_missing_census_bands_not_filled_with_zero(self):
  data=json.loads((ROOT/'negocio/05-Precios-y-Finanzas/universo-census.json').read_text())
  narrow=next(x for x in data['sectors'] if x['naics']=='333112')
  self.assertFalse(narrow['complete_target_bands']);self.assertEqual(len(narrow['missing_bands']),7)
 def test_tam_sam_arithmetic_and_units(self):
  d=m.sized(self.c,self.s)
  self.assertEqual(d['annual_price_assumed'],D(17988));self.assertEqual(d['tam_core_price_scenario'],D(53514300))
  self.assertEqual(d['sam_expected_accounts_scenario'],D('589.05'));self.assertEqual(d['sam_arr_scenario'],D('10595831.40'))
 def test_zero_conversion_produces_zero_customers_not_negative_cash_magic(self):
  s={**self.s,'paid_rate':D(0)};rows=m.forecast(self.c,s)
  self.assertTrue(all(r['active_logos_eom']==0 for r in rows));self.assertTrue(all(r['recurring_revenue']==0 for r in rows));self.assertLess(rows[-1]['cash_before_financing'],0)
 def test_no_sales_before_lag_and_outreach_start(self):
  rows=m.forecast(self.c,self.s)
  self.assertTrue(all(r['new_logos']==0 for r in rows[:4]));self.assertEqual(rows[4]['new_logos'],D('1.05'))
 def test_funnel_outreach_is_unique_and_bounded(self):
  for s in self.c['scenarios'].values():
   rows=m.forecast(self.c,s)
   self.assertLessEqual(sum(r['new_accounts_targeted'] for r in rows),self.c['core_firms'])
   self.assertLessEqual(rows[-1]['active_logos_eom'],m.sized(self.c,s)['sam_expected_accounts_scenario'])
 def test_recurring_revenue_not_exit_arr(self):
  rows=m.forecast(self.c,self.s);self.assertNotEqual(sum(r['recurring_revenue'] for r in rows[:12]),rows[11]['arr_exit'])
 def test_account_rollforward_and_queue_conservation(self):
  for s in self.c['scenarios'].values():
   active=D(0);backlog=D(0)
   for r in m.forecast(self.c,s):
    self.assertEqual(r['active_logos_eom'],active-r['lost_logos']+r['new_logos'])
    self.assertEqual(r['backlog_eom'],backlog+r['matured_wins']-r['new_logos'])
    self.assertGreaterEqual(r['backlog_eom'],0);active=r['active_logos_eom'];backlog=r['backlog_eom']
 def test_cash_rollforward(self):
  cash=self.c['forecast']['initial_cash']
  for r in m.forecast(self.c,self.s):
   cash+=r['net_cash_flow_before_financing'];self.assertEqual(cash,r['cash_before_financing'])
 def test_cogs_token_prices_correct_and_retry_not_free(self):
  u=m.monthly_cost(self.c,1499)
  self.assertEqual(u['extraction'],D('53.75'));self.assertEqual(u['escalation'],D('27.5'));self.assertEqual(u['recurring_cogs_per_active_month'],D('270.9025'))
 def test_revenue_doubling_does_not_alter_client_count(self):
  s={**self.s,'monthly_price':self.s['monthly_price']*2};a=m.forecast(self.c,self.s);b=m.forecast(self.c,s)
  self.assertEqual(a[-1]['active_logos_eom'],b[-1]['active_logos_eom']);self.assertEqual(a[-1]['arr_exit']*2,b[-1]['arr_exit'])
 def test_workbook_has_values_not_uncomputed_formula_promises(self):
  import openpyxl
  wb=openpyxl.load_workbook(ROOT/'negocio/05-Precios-y-Finanzas/VEXA-MERCADO-Y-FINANZAS.xlsx',read_only=True,data_only=True)
  self.assertEqual(wb['Forecast36'].max_row,109);self.assertEqual(wb['Resumen'].max_row,10);wb.close()
if __name__=='__main__':unittest.main()
